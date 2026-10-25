"""Run with the Codex bundled document Python runtime; no application imports required."""
import json,sys,platform,subprocess,xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'output/docx';OUT.mkdir(parents=True,exist_ok=True)
NOW=datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
def read(name,default=None):
 p=ROOT/name
 return json.loads(p.read_text()) if p.exists() else default
req=read('docs/requirements.json',[]);dec=read('docs/decisions.json',[])

def base(title,summary):
 d=Document();s=d.sections[0];s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
 for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.space_after=Pt(6)
 d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.line_spacing=1.12
 d.styles['Title'].font.size=Pt(25);d.styles['Title'].font.bold=True
 for n,size in [('Heading 1',16),('Heading 2',12),('Heading 3',10)]:d.styles[n].font.size=Pt(size)
 d.add_paragraph(title,'Title');d.add_paragraph('Miyazaki Industrial AI Copilot','Subtitle');d.add_paragraph(summary)
 d.add_paragraph('Updated '+NOW+' | Build 0.1.0 | Audience Project owner engineering and deployment reviewers')
 return d

def p(d,text,boldlead=None):
 para=d.add_paragraph()
 if boldlead:para.add_run(boldlead+' ').bold=True
 para.add_run(str(text));return para

def table(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
 trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 for c in t.rows[0].cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEFEA');c._tc.get_or_add_tcPr().append(sh)
  for run in c.paragraphs[0].runs:run.bold=True;run.font.size=Pt(9)
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):
   cells[i].text=str(value)
   for para in cells[i].paragraphs:
    para.paragraph_format.space_after=Pt(4)
    for run in para.runs:run.font.size=Pt(9)
  pr=cells[0]._tc.getparent().get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
 return t

# Actual test records; assertions are not expanded into claims about untested acceptance criteria.
testcases=[];jp=ROOT/'evidence/all-junit.xml'
if jp.exists():
 for x in ET.parse(jp).iter('testcase'):
  status='Failed' if x.find('failure') is not None or x.find('error') is not None else 'Not run' if x.find('skipped') is not None else 'Passed'
  testcases.append({'name':x.attrib['name'],'status':status,'seconds':x.attrib.get('time',''),'kind':'pytest','source':'evidence/all-junit.xml','command':'python -m pytest tests -q --junitxml=evidence/all-junit.xml','expected':'All assertions in '+x.attrib['name']+' hold against the isolated PostgreSQL/pgvector database.','actual':'All assertions passed.' if status=='Passed' else 'See JUnit failure or skip detail.'})
for filename,kind in [('evidence/browser-results.json','Browser'),('evidence/live-model-results.json','Live model')]:
 for x in (read(filename,{}) or {}).get('results',[]):testcases.append({**x,'kind':kind,'source':filename,'command':'node scripts/browser-tests.mjs' if kind=='Browser' else 'python scripts/live_smoke.py','expected':'Scenario assertions complete successfully with no substituted model response.','actual':x.get('actual','Scenario assertions passed.'),'environment':'evidence/environment.json'})
(ROOT/'evidence/test-index.json').write_text(json.dumps(testcases,ensure_ascii=False,indent=2))

first=base('Miyazaki Implementation Decisions and Change Log','The current build runs locally with a React interface, FastAPI services, PostgreSQL with pgvector and local Ollama inference. This document records decisions, implementation changes, evidence and limitations. It is not a production certification.')
first.add_heading('Architecture and setup',level=1)
table(first,['Subsystem','Implementation'],[['Frontend','React TypeScript Vite Tailwind TipTap paired Japanese and English resources'],['Application','FastAPI and SQLAlchemy in one application process'],['Database','PostgreSQL and pgvector with exact workspace filtered cosine search'],['Jobs','PostgreSQL queue and two in process worker lanes with APScheduler'],['Models','Configurable Ollama generation and multilingual embeddings'],['Files','Workspace partitioned local files and Japanese font DOCX PDF Markdown CSV ZIP exports']])
p(first,'Create a Python virtual environment, install requirements.txt, copy .env.example to .env, configure PostgreSQL and run the commands below. pgvector must be installed for the PostgreSQL major version.')
for cmd in ['python manage.py setup','python manage.py check','python manage.py run','python manage.py seed --reset']:p(first,cmd)
p(first,'Node is required to rebuild the frontend, but built assets are shipped for demo machines. setup is idempotent; reset is restricted to demo configuration. Local service installation remains the responsibility of the machine owner.')
first.add_heading('Decisions and consequences',level=1)
for item in dec:
 first.add_heading(item['id']+' '+item['decision'],level=2)
 p(first,item['why'],'Reason');p(first,item['alternative'],'Alternative');p(first,item['consequence'],'Consequence')
first.add_heading('Implementation changes',level=1)
changes=read('docs/changes.json',[])
for change in changes:
 first.add_heading(change['area'],level=2)
 p(first,change['date']+' | '+', '.join(change['requirements']),'Date and requirements')
 p(first,change['before'],'Before');p(first,change['after'],'Changed');p(first,change['why'],'Why');p(first,change['evidence'],'Evidence')
first.add_heading('Verification snapshot',level=1)
counts=Counter(t['status'] for t in testcases);p(first,'Recorded cases '+str(dict(counts))+'. Detailed scenario results are in the separate coverage and security report. A pass proves only the recorded assertions.')
env=read('evidence/environment.json',{});table(first,['Environment','Recorded value'],[[k,str(v)] for k,v in env.items() if k!='packages'])
pre=read('evidence/preflight.json',{});p(first,json.dumps({k:v for k,v in pre.items() if k not in ['generation']},ensure_ascii=False),'Preflight')
first.add_heading('Known limitations and production follow up',level=1)
for text in ['Japanese OCR and scanned image interpretation are explicitly deferred.','Draft Atomic Loops learning material Japanese wording and legal notices require external approval.','Some included stories remain partial. The accompanying 98 story ledger identifies missing acceptance criteria.','Local demo MFA bypass and development outbox are unsuitable for shared network or production use.','No independent penetration test load test accessibility audit production hosting review or backup restore SLA has been established.','AI citation validity is not equivalent to factual entailment. Generated claims require review and the full SRS golden set has not passed.']:
 p(first,text)
first.add_heading('Evidence and source documents',level=1)
for s in ['Miyazaki Industrial AI Copilot — SRS.pdf sections 4 6 9 and 11','Codex Build Brief All MUSTs, Minimal Setup.pdf','User overrides pgvector required and Japanese OCR deferred','docs/requirements.json docs/decisions.json evidence/test-index.json']:p(first,s)
for para in first.paragraphs:
 if para.style.name.startswith('Heading'):para.text=__import__('re').sub(r'[^\w\s]',' ',para.text)
first.save(OUT/'Miyazaki Implementation Decisions and Change Log.docx')

second=base('Miyazaki Requirements Coverage and Security Verification','This report separates required scope from actual implementation and observed test results. All 98 SRS stories and all 20 non functional requirements are recorded. Included does not mean complete and a passing subset test does not establish full SRS compliance.')
second.add_heading('Readiness conclusion',level=1)
p(second,'Controlled localhost demonstration only. Shared network and production deployment remain blocked pending hardening, content approval and closure of incomplete acceptance criteria. See the requirement entries for exact gaps.')
stats=Counter(r['status'] for r in req if r['included']);table(second,['Measure','Count'],[['SRS stories',len(req)],['Must stories',sum(r['priority']=='Must' for r in req)],['Included stories including supporting scope',sum(r['included'] for r in req)],['Included implemented',stats['Implemented']],['Included partial',stats['Partial']],['Deferred',sum(not r['included'] for r in req)]])
p(second,'The base target is 67 Must stories plus six named v1 Should stories. Supporting inclusions are K07 F04 and the E02 required D09 conversion path. OCR portions of E05 F01 and K04 are outside this phase by explicit user instruction.')
second.add_heading('Feature inventory and access',level=1)
table(second,['Area','Routes and access'],[['Authentication','/login /invite/{token} /reset authenticated sessions and MFA'],['Operator','/admin companies shared company record feedback audit outbox'],['Platform editor','/admin/library bilingual draft review publish'],['Company','/w/{company}/home profile workflows chats knowledge ai-use-cases feedback settings'],['Advisor','Granted shared profile read or comment only'],['Health','/health and python manage.py check local readiness']])
second.add_heading('Requirements by SRS identifier',level=1)
current=''
for r in req:
 epic=r['id'][3]
 if epic!=current:second.add_heading('Epic '+epic,level=2);current=epic
 second.add_heading(r['id']+' '+r['priority']+' '+r['status'],level=3)
 p(second,r['story'])
 p(second,r['acceptance_criteria'] or 'See the original SRS section 4 for the complete acceptance text.','Acceptance')
 p(second,('Included' if r['included'] else 'Deferred')+' — '+r['scope_reason'],'Scope')
 p(second,r['behaviour'],'Implemented behaviour')
 if r['gaps']:p(second,r['gaps'],'Remaining gaps')
 matched=[]
 for ref in r['tests']:
  found=[t for t in testcases if ref.replace('browser ','').replace('live ','') in t['name']]
  matched+=found
 outcome='Passed subset only' if matched and all(t['status']=='Passed' for t in matched) else 'Failed or blocked subset' if any(t['status']=='Failed' for t in matched) else 'Not run or no directly mapped assertion'
 p(second,outcome+'. '+('; '.join(r['tests']) or 'No tests for deferred functionality.'),'Verification')
second.add_heading('Non functional requirements',level=1)
nfr=[('01','Streaming and page timings measured','Two second p95 not established'),('02','Job progress caching and local generation','200 pages in five minutes and 2000 words in 90 seconds unverified'),('03','Single application instance','500 workspaces 5000 users 200 concurrent and horizontal scale untested'),('04','Health endpoint graceful failure','99.5 percent SLA and maintenance operations not established'),('05','Portable data structure and backup instructions','Automated backups RPO RTO and quarterly restore test not established'),('06','Hashing validation security headers dependency audits','ASVS assessment TLS at rest encryption vault and independent penetration test outstanding'),('07','Server roles MFA sessions','Demo bypass enabled intentionally; full security review outstanding'),('08','Scoped API and vector tests','Broader adversarial corpus and continuous delivery pipeline outstanding'),('09','Insert only database audit and histories','One year retention operations and independent tamper review outstanding'),('10','Consent gates local redaction and minimal logs','APPI legal review and 30 day request operations outstanding'),('11','Local inference and storage','Approved Japan hosting provider agreements and processing regions outstanding'),('12','Retention configuration ZIP and deletion request','Complete approved deletion and backup expiry operations outstanding'),('13','Bilingual resources and Japanese fonts','Some English labels and Japanese human approval outstanding'),('14','Keyboard labels responsive layouts','WCAG and JIS audit not performed'),('15','Tour and responsive core flow','One hour user study not performed'),('16','Provenance citations markers injection fixtures','Full factual extraction citation coverage recommendation and adversarial golden set not passed'),('17','Model metadata usage jobs and errors','Distributed tracing metrics dashboards alerts outstanding'),('18','Versioned resources migrations and tests','Deferred content consoles and feature flags outstanding'),('19','Chromium desktop and mobile viewport checks','Other browsers operating systems and physical devices untested'),('20','Atomic budgets and usage exports','Hosted provider costs and report presentation need completion')]
table(second,['NFR','Present','Remaining verification'],[['NFR-'+a,b,c] for a,b,c in nfr])
second.add_heading('Executed tests',level=1)
p(second,'Environment Python application virtual environment local PostgreSQL with pgvector and Ollama. Backend tests use a separate security demo database. Mocked provider tests identify the mock in their name; live model cases use real local services.')
for kind in ['pytest','Browser','Live model']:
 group=[t for t in testcases if t['kind']==kind]
 second.add_heading(kind+' results',level=2)
 table(second,['Scenario','Result','Actual observation'],[[t['name'],t['status'],t.get('actual','Assertions passed.')] for t in group])
 for t in group:
  if t['status']!='Passed':p(second,t.get('actual','No detail')[:800],t['name'])
second.add_heading('Security coverage and concerns',level=1)
security=[('Isolation','API IDs vectors source files jobs links exports and caches','Passing subset tests; no independent penetration test'),('Authentication','CSRF cookies session expiry logout generic login/reset rate limits MFA recovery replay','TLS secure cookies and privileged secret protection remain deployment tasks'),('Untrusted input','SQL identifier injection editor HTML file extension and ZIP checks path traversal','Full parser fuzzing resource exhaustion and malicious office corpus untested'),('Crawling','Private loopback link local destinations redirects pinned addresses robots limits','Live site corpus DNS edge cases and production egress rules require validation'),('AI','Quarantine policy context toggles rejected claims citation handle validity and confirmations','Semantic truth and every statement provenance are not guaranteed by handle checks'),('Privacy','Outbox authorisation consent evaluation redaction and secret free error design','Redaction review required; legal privacy notices remain draft'),('Concurrency','Budget reservation action replay job leases and digest keys','Load stress crash mid transaction and multi process operation not validated')]
table(second,['Area','Checks','Concern'],security)
second.add_heading('Dependency advisories',level=1)
pa=read('evidence/python-audit.json',{});fa=read('evidence/frontend-audit.json',{})
p(second,'Python vulnerabilities currently reported '+str(sum(len(x.get('vulns',[])) for x in pa.get('dependencies',[])))+'. Frontend '+json.dumps(fa.get('metadata',{}).get('vulnerabilities',{}))+'. Advisory results are time specific and not a complete security assessment.')
second.add_heading('Production risk register',level=1)
for risk in read('docs/risks.json',[]):
 second.add_heading(risk['id']+' '+risk['area'],level=2)
 p(second,risk['severity']+' | '+risk['status'],'Assessment');p(second,risk['concern'],'Concern');p(second,risk['required'],'Required before production')
second.add_heading('Deployment blockers and approvals',level=1)
for text in ['Keep the current demo bound to localhost with synthetic data.','Before shared network use remove demo credentials and MFA bypass disable development outbox configure TLS trusted hosts secure cookies and appropriate database privileges.','Before production complete included acceptance gaps independently test security accessibility backup restore retention deletion performance and provider failure behaviour.','Obtain Operator approval for legal notices Japanese wording data controller hosting region and any hosted model processing.','Review every AI generated draft. Source handles prove traceability only; factual correctness still requires evidence and human review.']:
 p(second,text)
second.add_heading('Evidence locations',level=1)
for location in ['docs/requirements.json','docs/decisions.json','evidence/all-junit.xml','evidence/browser-results.json','evidence/live-model-results.json','evidence/preflight.json','evidence/python-audit.json','evidence/frontend-audit.json','evidence/test-index.json']:p(second,location)
for para in second.paragraphs:
 if para.style.name.startswith('Heading'):para.text=__import__('re').sub(r'[^\w\s]',' ',para.text)
second.save(OUT/'Miyazaki Requirements Coverage and Security Verification.docx')
print('Wrote two DOCX reports',NOW)
