import re, json, zipfile, html
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from .config import ROOT

FONT=ROOT/'frontend/public/fonts/NotoSansJP-Regular.ttf'

def export_text(title,content,fmt,markers=True):
    if not markers:content=re.sub(r'\[(?:F|S|M):[^\]]+\]|\(AI assumption, verify\)|AIの仮定・要確認','',content)
    if fmt=='md':return ('# '+title+'\n\n'+content).encode('utf-8'),'text/markdown','.md'
    out=BytesIO()
    if fmt=='docx':
        d=Document();normal=d.styles['Normal'];normal.font.name='Noto Sans JP';normal.font.size=Pt(10)
        normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans JP')
        d.add_heading(title,0)
        for line in content.splitlines():
            if line.startswith('#'):d.add_heading(line.lstrip('# ').strip(),min(len(line)-len(line.lstrip('#')),3))
            else:d.add_paragraph(line)
        d.save(out);return out.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document','.docx'
    if fmt=='pdf':
        if not FONT.exists():raise ValueError('Bundled Japanese TrueType font is missing. Rebuild the asset package.')
        if 'NotoJP' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('NotoJP',str(FONT)))
        style=ParagraphStyle('Japanese',fontName='NotoJP',fontSize=9,leading=15,wordWrap='CJK',spaceAfter=7)
        heading=ParagraphStyle('JapaneseHeading',parent=style,fontSize=17,leading=24,spaceAfter=16)
        story=[Paragraph(html.escape(title),heading)]
        for line in content.splitlines():story.append(Paragraph(html.escape(line) or ' ',style))
        SimpleDocTemplate(out,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42).build(story)
        return out.getvalue(),'application/pdf','.pdf'
    raise ValueError('Choose docx, pdf or md.')
