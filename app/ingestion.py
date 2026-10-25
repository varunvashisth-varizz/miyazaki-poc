import re, socket, ipaddress, urllib.robotparser
from urllib.parse import urlsplit, urljoin
from pathlib import Path
from bs4 import BeautifulSoup
import httpx
from .config import ROOT
from .security import public_url, archive_safe

SUSPICIOUS=re.compile(r'ignore\s+(all\s+)?(previous|prior)\s+instructions|system\s*:|reveal\s+.*secret|以前の指示を無視',re.I)

def html_text(content):
    soup=BeautifulSoup(content,'html.parser')
    for t in soup(['script','style','noscript','iframe','svg']):t.decompose()
    text=soup.get_text(' ',strip=True)
    flagged=bool(SUSPICIOUS.search(text))
    # Quarantine the entire source; regex detection is not the sole policy boundary.
    return text,flagged

def fetch(url,maximum=2*1024*1024):
    """Resolve then pin each redirect hop; never follow redirects automatically."""
    for _ in range(5):
        public_url(url);parts=urlsplit(url);host=parts.hostname
        addresses=socket.getaddrinfo(host,parts.port or (443 if parts.scheme=='https' else 80),type=socket.SOCK_STREAM)
        if any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):raise ValueError('Unsafe DNS destination.')
        ip=addresses[0][4][0];netloc='['+ip+']' if ':' in ip else ip
        pinned=parts._replace(netloc=netloc).geturl()
        with httpx.Client(timeout=15,trust_env=False) as c:
            with c.stream('GET',pinned,headers={'Host':parts.netloc,'User-Agent':'MiyazakiCopilot/0.1'},extensions={'sni_hostname':host.encode()}) as r:
                if r.status_code in (301,302,303,307,308):url=urljoin(url,r.headers.get('location',''));continue
                r.raise_for_status()
                if int(r.headers.get('content-length','0'))>maximum:raise ValueError('Page exceeds size limit.')
                result=bytearray()
                for data in r.iter_bytes():
                    result.extend(data)
                    if len(result)>maximum:raise ValueError('Page exceeds size limit.')
                return bytes(result),url
    raise ValueError('Too many redirects.')

def crawl(url,limit=20,depth=2,fixture=False,injection=False,progress=None):
    if fixture:
        file=ROOT/'fixtures'/('injection.html' if injection else 'company.html')
        text,flagged=html_text(file.read_bytes())
        if progress:progress(1,'Read saved offline snapshot')
        return [{'text':text,'url':'fixture://'+file.name,'flagged':flagged,'snapshot':True}]
    public_url(url);parts=urlsplit(url);root=parts.scheme+'://'+parts.netloc
    robots=urllib.robotparser.RobotFileParser();body,_=fetch(root+'/robots.txt');robots.parse(body.decode('utf-8','replace').splitlines())
    queue=[(url,0)];seen=set();pages=[]
    while queue and len(pages)<min(limit,200):
        current,level=queue.pop(0)
        if current in seen:continue
        seen.add(current)
        if urlsplit(current).netloc!=parts.netloc or not robots.can_fetch('MiyazakiCopilot',current):continue
        raw,final=fetch(current)
        if urlsplit(final).netloc!=parts.netloc:continue
        text,flagged=html_text(raw);pages.append({'text':text,'url':final,'flagged':flagged,'snapshot':False})
        if progress:progress(len(pages),'Pages read: '+str(len(pages)))
        if level<depth:
            soup=BeautifulSoup(raw,'html.parser')
            for a in soup.find_all('a',href=True):
                link=urljoin(final,a['href']).split('#')[0]
                if urlsplit(link).netloc==parts.netloc and link not in seen:queue.append((link,level+1))
    if not pages:raise ValueError('No readable pages. Check robots rules or describe your company manually.')
    return pages

def parse_file(path):
    ext=path.suffix.lower();pieces=[]
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if len(reader.pages)>500:raise ValueError('PDF exceeds 500-page processing limit.')
        for i,p in enumerate(reader.pages):
            text=p.extract_text() or ''
            if text.strip():pieces.append((f'Page {i+1}',text))
        if not pieces:raise ValueError('Scanned PDF: OCR is deferred. Upload selectable text or enter the content manually.')
    elif ext in ('.docx','.xlsx','.pptx'):
        archive_safe(path)
        if ext=='.docx':
            from docx import Document
            d=Document(path);pieces=[('Document','\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows))]
        elif ext=='.xlsx':
            from openpyxl import load_workbook
            wb=load_workbook(path,read_only=True,data_only=True)
            for sh in wb:
                if sh.max_row>50000 or sh.max_column>200:raise ValueError('Spreadsheet exceeds safe processing limits.')
                pieces.append((sh.title,'\n'.join(' | '.join(str(v or '') for v in row) for row in sh.iter_rows(values_only=True))))
            wb.close()
        else:
            from pptx import Presentation
            p=Presentation(path)
            if len(p.slides)>500:raise ValueError('Presentation exceeds slide limit.')
            for i,slide in enumerate(p.slides):
                content=[]
                for sh in slide.shapes:
                    if sh.has_text_frame:content.append(sh.text)
                    if sh.has_table:content.extend(' | '.join(c.text for c in row.cells) for row in sh.table.rows)
                pieces.append((f'Slide {i+1}','\n'.join(content)))
    elif ext in ('.png','.jpg','.jpeg','.webp'):
        from PIL import Image
        with Image.open(path) as im:im.verify()
        raise ValueError('Image stored for preview. OCR and visual interpretation are deferred; add a manual description.')
    elif ext in ('.txt','.md','.csv'):
        pieces=[('Text',path.read_text(encoding='utf-8-sig'))]
    else:raise ValueError('Unsupported file type.')
    if sum(len(t) for _,t in pieces)>4_000_000:raise ValueError('Extracted content exceeds processing limit.')
    return pieces

def chunk_text(text,size=1000,overlap=120):
    return [text[i:i+size] for i in range(0,len(text),size-overlap) if text[i:i+size].strip()]
